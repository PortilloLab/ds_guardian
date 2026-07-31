import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_docx():
    doc = docx.Document()
    
    # Page setup
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Arial'
    font.size = Pt(11)
    
    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("INFORME DE TRANSFERENCIA Y PROMPT DE GENERACIÓN PARA IA\nPROYECTO: FRAMEWORK 'DS GUARDIAN'")
    run.bold = True
    run.font.size = Pt(16)
    
    doc.add_paragraph("Este documento contiene las especificaciones técnicas y el prompt exacto que debes proveer a otra Inteligencia Artificial (IA) en tu otra computadora para que genere y replique exactamente el framework DS Guardian robusto de Ciencia de Datos.").paragraph_format.space_after = Pt(12)
    
    # Heading 1: Prompt
    h1 = doc.add_paragraph()
    run = h1.add_run("1. Prompt para copiar y pegar en la otra IA")
    run.bold = True
    run.font.size = Pt(14)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(6)
    
    p_prompt = doc.add_paragraph()
    p_prompt.paragraph_format.left_indent = Inches(0.25)
    run_prompt = p_prompt.add_run(
        "\"\"\"\n"
        "Actúa como un experto en Ciencia de Datos y Machine Learning en Python.\n"
        "Quiero que generes una librería local llamada 'ds_guardian' compuesta por módulos robustos para el procesamiento, modelado, visualización y auditoría avanzada de datos de proyectos reales. La librería debe tener la siguiente estructura y cumplir con los requisitos indicados para cada archivo:\n\n"
        "Estructura de la carpeta:\n"
        "ds_guardian/\n"
        "├── __init__.py\n"
        "├── eda.py\n"
        "├── limpieza.py\n"
        "├── visualizacion.py\n"
        "├── modelos.py\n"
        "└── auditoria.py\n\n"
        "--- ESPECIFICACIONES DE LOS ARCHIVOS ---\n\n"
        "1. __init__.py:\n"
        "Debe importar todas las funciones de los módulos internos para permitir importaciones rápidas (ej: 'from ds_guardian import eda, limpieza, modelos, auditoria, visualizacion').\n\n"
        "2. eda.py (Análisis Exploratorio y Tratamiento de Outliers):\n"
        "- resumir_datos(df): Imprime dimensiones, dtypes y estadísticas descriptivas de forma segura.\n"
        "- missing_values_table(df): Muestra el conteo y porcentaje de nulos por columna.\n"
        "- optimizar_memoria(df): Reduce tipos de datos de columnas numéricas verificando 'pd.api.types.is_numeric_dtype(df[col])' para evitar fallar con strings.\n"
        "- detectar_outliers_iqr(df): Encuentra y reporta outliers en variables numéricas usando el rango intercuartílico (IQR).\n"
        "- acotar_outliers_iqr(df, columnas=None): Aplica Winsorization/Capping acotando valores extremos dentro de los límites del IQR sin eliminar filas.\n\n"
        "3. limpieza.py (Preprocesamiento sin Data Leakage):\n"
        "- imputar_nulos(df_train, df_test=None, estrategia_num='median', estrategia_cat='most_frequent'): Ajusta SimpleImputer solo en Train y transforma Train y Test por separado.\n"
        "- tratar_duplicados(df): Identifica y elimina filas duplicadas.\n"
        "- codificar_variables(df_train, df_test=None): Aplica One-Hot Encoding alineando columnas entre Train y Test.\n"
        "- escalar_caracteristicas(df_train, df_test=None, metodo='standard', columnas=None): Escala características numéricas usando StandardScaler o MinMaxScaler ajustando solo en Train.\n\n"
        "4. visualizacion.py (Gráficos premium):\n"
        "- configurar_estilo(theme='light'): Configura el estilo de seaborn/matplotlib. Soporta tema 'light' y 'dark' premium con paletas ajustadas (muted o mako).\n"
        "- plot_distribucion(df, columna): Histograma con KDE controlando nulos.\n"
        "- plot_categorica(df, columna, max_categorias=20): Gráfico de barras horizontales.\n"
        "- plot_correlacion(df): Heatmap de correlación solo sobre variables numéricas.\n"
        "- plot_importancia_caracteristicas(modelo, feature_names, max_features=15): Grafica la importancia de variables para modelos de árboles.\n\n"
        "5. modelos.py (Modelado y Optimización):\n"
        "- evaluar_clasificacion(y_true, y_pred, y_prob=None): Accuracy, Reporte, Matriz de Confusión y ROC-AUC (binario y multiclase).\n"
        "- evaluar_regresion(y_true, y_pred): MSE, RMSE, MAE y R2 Score sin usar parámetros obsoletos.\n"
        "- validacion_cruzada(modelo, X, y, cv=5, scoring='accuracy'): K-Fold cross validation detallado.\n"
        "- optimizar_hiperparametros(modelo, param_grid, X, y, cv=3, n_iter=10, scoring='accuracy'): Optimiza hiperparámetros usando RandomizedSearchCV y retorna el mejor modelo.\n\n"
        "6. auditoria.py (QA Auditor Avanzado):\n"
        "- revisar_datos_finales(df, y=None): Realiza un checklist completo. Comprueba nulos, dtypes correctos, desbalanceo de clases (si max_pct > 65% en y), multicolinealidad (correlación entre columnas > 0.95) y Data Leakage (correlación de features con y > 0.99). Imprime salidas con colores ANSI en consola.\n"
        "- registrar_y_comparar_modelo(nombre_proyecto, metrica, valor, maximizar=True): Guarda métricas en 'historial_proyectos.json' y las compara alertando con colores en consola si mejoró o empeoró.\n\n"
        "Genera todo el código completo y documentado para cada archivo de forma modular y profesional.\n"
        "\"\"\""
    )
    run_prompt.italic = True
    run_prompt.font.size = Pt(10)
    
    # Heading 2: Arquitectura del Skill
    h2 = doc.add_paragraph()
    run = h2.add_run("2. Detalles Críticos y Arquitectura")
    run.bold = True
    run.font.size = Pt(14)
    h2.paragraph_format.space_before = Pt(18)
    h2.paragraph_format.space_after = Pt(6)
    
    bullets = [
        ("Evitar Data Leakage:", " Toda imputación, codificación y escalamiento debe ajustarse únicamente en Train y aplicarse por separado a Test."),
        ("Detección de Fugas e Imbalances:", " El auditor integrado escanea desbalances en el target y correlaciones sospechosas (>0.99) que sugieren fugas de información."),
        ("Tuning de Modelos Integrado:", " El módulo de modelos ahora cuenta con RandomizedSearchCV integrado de forma nativa para sintonizar los modelos antes de producción."),
        ("Winsorization (Capping):", " En lugar de descartar datos de outliers, el framework permite acotar extremos usando los límites del IQR.")
    ]
    
    for title_b, desc_b in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        r_title = bp.add_run(title_b)
        r_title.bold = True
        bp.add_run(desc_b)
        
    doc.save("Instrucciones_Generacion_AI.docx")

if __name__ == "__main__":
    build_docx()
